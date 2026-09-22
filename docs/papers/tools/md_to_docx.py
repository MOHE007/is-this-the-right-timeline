#!/usr/bin/env python3
"""把论文 Markdown 转成排版好的 Word（.docx）。

用法:
    python3 md_to_docx.py <input.md> <output.docx>

支持本项目论文用到的元素：YAML 头、多级标题、管道表格、有序/无序列表、
代码块、引用块、水平线、粗体与行内代码。中文字体显式设置 eastAsia，
避免 Word/WPS 打开时中文变成默认字体。
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_FONT_EA = "宋体"
HEAD_FONT_EA = "黑体"
BODY_FONT_LATIN = "Times New Roman"
CODE_FONT = "Consolas"
ACCENT = RGBColor(0x8A, 0x3A, 0x26)  # 与游戏内的朱砂色一致


def set_run_font(run, ea=BODY_FONT_EA, latin=BODY_FONT_LATIN, size=None, bold=None, color=None):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), ea)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_inline(paragraph, text, base_size=10.5, ea=BODY_FONT_EA):
    """处理 **粗体** 与 `行内代码`。"""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, ea=ea, size=base_size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, ea=ea, latin=CODE_FONT, size=base_size - 0.5)
            run.font.color.rgb = ACCENT
        else:
            run = paragraph.add_run(part)
            set_run_font(run, ea=ea, size=base_size)


def shade(cell, hex_color="F2F0EA"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        add_inline(cell.paragraphs[0], text, base_size=10, ea=HEAD_FONT_EA)
        for run in cell.paragraphs[0].runs:
            run.bold = True
        shade(cell)
    for row in body:
        cells = table.add_row().cells
        for i, text in enumerate(row[: len(header)]):
            cells[i].text = ""
            add_inline(cells[i].paragraphs[0], text, base_size=10)
    doc.add_paragraph()


def convert(md_path: Path, docx_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").split("\n")

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)  # A4
    for attr in ("top_margin", "bottom_margin"):
        setattr(section, attr, Cm(2.2))
    for attr in ("left_margin", "right_margin"):
        setattr(section, attr, Cm(2.0))

    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.font.name = BODY_FONT_LATIN
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT_EA)

    i = 0
    in_code = False
    code_buffer: list[str] = []
    in_frontmatter = False
    table_buffer: list[list[str]] = []
    last_heading = [0]  # 上一个标题的层级（用于识别 h1 后的副标题）

    def flush_table():
        nonlocal table_buffer
        if table_buffer:
            add_table(doc, table_buffer)
            table_buffer = []

    while i < len(lines):
        line = lines[i].rstrip()

        # YAML 头：跳过
        if i == 0 and line.strip() == "---":
            in_frontmatter = True
            i += 1
            continue
        if in_frontmatter:
            if line.strip() == "---":
                in_frontmatter = False
            i += 1
            continue

        # 代码块
        if line.startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.5)
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(8)
                run = p.add_run("\n".join(code_buffer))
                set_run_font(run, ea=BODY_FONT_EA, latin=CODE_FONT, size=9)
                code_buffer = []
                in_code = False
            else:
                flush_table()
                in_code = True
            i += 1
            continue
        if in_code:
            code_buffer.append(line)
            i += 1
            continue

        # 表格
        if line.startswith("|") and line.endswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                i += 1
                continue
            table_buffer.append(cells)
            i += 1
            continue
        flush_table()

        stripped = line.strip()

        # 空行 / 水平线
        if not stripped:
            i += 1
            continue
        if re.fullmatch(r"-{3,}", stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            pPr = p._p.get_or_add_pPr()
            bdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "BBBBBB")
            bdr.append(bottom)
            pPr.append(bdr)
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level, text = len(m.group(1)), m.group(2)
            sizes = {1: 20, 2: 15, 3: 12.5, 4: 11.5}
            # 紧跟在 h1 之后的 h3 是副标题（"——以……为例"）：它不是章节标题，
            # 若按 Heading 3 输出会以"3 级标题"混进 Word 导航窗格，故单独居中排版。
            if level == 3 and last_heading[0] == 1:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(18)
                run = p.add_run(text)
                set_run_font(run, ea=HEAD_FONT_EA, size=11.5, bold=False,
                             color=RGBColor(0x6B, 0x62, 0x59))
                last_heading[0] = level
                i += 1
                continue
            # 用真正的 Heading 样式：Word 才能生成导航窗格与大纲结构
            # （普通段落做标题会让长篇文档无法跳转）。字体与配色在 run 层
            # 显式覆盖，保证与 PDF 的观感一致。
            p = doc.add_paragraph(style=f"Heading {min(level, 4)}")
            p.paragraph_format.space_before = Pt(14 if level <= 2 else 10)
            p.paragraph_format.space_after = Pt(6)
            if level == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(text)
            set_run_font(run, ea=HEAD_FONT_EA, size=sizes.get(level, 11), bold=True,
                         color=ACCENT if level <= 2 else RGBColor(0x3A, 0x37, 0x33))
            last_heading[0] = level
            i += 1
            continue

        # 引用块
        if stripped.startswith(">"):
            text = stripped.lstrip("> ").strip()
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            add_inline(p, text)
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(0x44, 0x44, 0x44)
            i += 1
            continue

        # 列表
        m = re.match(r"^([-*])\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, m.group(2))
            i += 1
            continue
        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            # 不使用 "List Number" 样式：python-docx 默认模板中该样式的所有段落
            # 共用同一个 style 级 numId，Word 会**跨节连续编号**——本文 §3.5 的
            # 1-3 之后，§7.1 会接着从 4 开始、§8.3 从 16 开始，而论文每一节都
            # 必须从 1 重新计数。因此改为「字面序号 + 悬挂缩进」，编号完全由
            # Markdown 原文决定，不依赖 Word 的编号定义。
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.85)
            p.paragraph_format.first_line_indent = Cm(-0.85)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(f"{m.group(1)}. ")
            set_run_font(run, ea=HEAD_FONT_EA, size=10.5)
            add_inline(p, m.group(2))
            i += 1
            continue

        # 正文
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(6)
        add_inline(p, stripped)
        i += 1

    flush_table()
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)
    print(f"wrote {docx_path}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        raise SystemExit(1)
    convert(Path(sys.argv[1]), Path(sys.argv[2]))
